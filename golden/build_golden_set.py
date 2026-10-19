"""Build golden_set.json and DWN_Golden_Data_Set_v2.xlsx from scenarios + reference engine."""
import io, json, os, sys
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reference_engine import ASSUMPTIONS, CLASSES, CLASS_NAMES, evaluate
from scenarios import CLIENTS, S

pct = lambda w: " / ".join(f"{k} {round(v*100)}" for k, v in w.items() if v or k != "SS")
rows, records = [], []
for s in S:
    c = CLIENTS[s["client"]]
    r = evaluate(c, s["event"])
    e = s["expected"]
    ok = (r["tier"] == e["tier"] and r["flags"] == e["flags"] and r["message"] == e["message"]
          and (not e["goal_after"] or r.get("goal_after") == e["goal_after"]))
    trade_txt = "none"
    if isinstance(r.get("trades"), dict):
        trade_txt = ", ".join(f"{'buy' if d > 0 else 'sell'} {CLASS_NAMES[k].lower()} {abs(d):.1f} pts" for k, d in r["trades"].items() if abs(d) >= 0.05)
    sig = s["event"].get("signal")
    records.append(dict(
        id=s["id"], category=s["category"], client=s["client"], title=s["title"], event=s["event_text"], event_input=s["event"],
        source=s["source"], expected=dict(tier=e["tier"], flags=sorted(e["flags"]), client_message=e["message"], goal_status_after=e["goal_after"]),
        must_include=s["must_include"], must_not_do=s["must_not"], must_not_trade_classes=s["must_not_trade"],
        expected_extraction=sig, tests=s["tests"],
        reference_result=dict(tier=r["tier"], flags=sorted(r["flags"]), client_message=r["message"], factors=r.get("factors"),
                              goal_before=r.get("goal_before"), goal_after=r.get("goal_after"), max_drift_pts=r.get("max_drift_pts"),
                              band_pts=r.get("band_pts"), proposed_adjustment=trade_txt, matches_expected=ok)))
    rows.append((s, c, r, e, ok, trade_txt, sig))

assert all(x[4] for x in rows), "reference engine disagrees with authored expectations"

with io.open(os.path.join(HERE, "golden_set.json"), "w", encoding="utf-8") as f:
    json.dump(dict(version="2.0", currency="USD", assumptions=ASSUMPTIONS,
                   clients={k: {kk: vv for kk, vv in v.items()} for k, v in CLIENTS.items()}, scenarios=records), f, indent=2, default=list)

# ---------------- workbook ----------------
wb = Workbook()
HDR = PatternFill("solid", fgColor="102844"); HF = Font(bold=True, color="FFFFFF"); WRAP = Alignment(wrap_text=True, vertical="top")
def sheet(name, header, data, widths):
    ws = wb.create_sheet(name)
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = HDR, HF, WRAP
    for d in data:
        ws.append(d)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "A2"
    return ws

ws = wb.active; ws.title = "About"
tiers = Counter(str(x[3]["tier"]) for x in rows)
about = [
 ("DWN Golden Data Set v2 (synthetic)", ""),
 ("Status", "DRAFT for advisor and compliance review. Replaces the earlier workbook, which was rejected by the product owner."),
 ("Currency / jurisdiction", "USD; US rules assumed (SEC fiduciary duty, Reg BI as benchmark, GDPR/CCPA for data). Per concept document section 12."),
 ("Contents", f"{len(S)} scenarios in 8 categories, 12 synthetic clients, tier framework, control checks, assumptions, reviewer sheet."),
 ("Expected tier mix", " / ".join(f"{k}: {tiers[k]}" for k in ("1", "2", "3", "Paused"))),
 ("How expectations were produced", "Tier, flags, message routing and goal status were authored per scenario from the concept-document tier framework (section 5). They were then checked against an independent reference implementation (golden/reference_engine.py)."),
 ("Honest calibration record", "On the first run 16 of 36 authored expectations matched the reference engine. The 20 mismatches came from: (a) goal targets not yet calibrated so that on-track baselines existed (most); (b) shock sizes that were too small to breach a band or too large and tripped the bad-tick limit (S-03, S-06, S-27, S-28); (c) expectations revised after review because a cash withdrawal or outflow legitimately moves cash outside the band (S-12, S-14, S-16); (d) a naming change: a client withdrawal request is flagged client_request, not confirmed_life_event. Goal targets and shock sizes were therefore tuned so each scenario produces its intended situation. (e) After the production engine was written, it disagreed on S-20: the reference engine had silently ignored a known $150,000 contribution when no note signal was present. The production behaviour (a known cash flow counts, so cash drifts out of band) is right, so the reference engine and the S-20 expectation were corrected. A 36/36 match shows the reference engine follows the framework as written; it does not show the framework is right."),
 ("Design decisions needing confirmation", "1) Size factor = share of portfolio moved by the proposed adjustment or by a confirmed cash flow (0 if no adjustment). The concept document says 'how much of the portfolio is affected' without defining it; measuring market loss instead would make almost every sell-off Tier 2. 2) Hidden instructions in a note => Paused (concept doc section 5), not Tier 1. 3) Conflicting sources, missing price and currency mismatch => Paused with an abstention message."),
 ("Limits", "All data is fictional. All thresholds and market assumptions are placeholders for AWM compliance and risk. Expected outputs are not yet reviewed by advisors or compliance. Draft wording is not covered here: it is tested separately for prohibited language and untraceable figures."),
 ("What is NOT tested here", "LLM text quality, live provider behaviour, UI, authentication, scheduling. These have their own tests."),
]
for a in about: ws.append(a)
ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 130
for row in ws.iter_rows():
    for c in row: c.alignment = WRAP
ws["A1"].font = Font(bold=True, size=14)
for r_ in range(2, ws.max_row + 1): ws.cell(r_, 1).font = Font(bold=True)

sheet("Golden Set",
 ["ID", "Category", "Scenario", "Client", "Client context (synthetic)", "Event / input", "Source", "Expected tier", "Required flags",
  "Expected client message", "Expected goal status after", "Must include", "Must not do", "Must not trade", "Expected note extraction", "What it tests",
  "Ref. tier", "Ref. message", "Max drift (pts)", "Band (pts)", "Goal before > after", "Proposed adjustment (direction, size)", "Match"],
 [[s["id"], s["category"], s["title"], s["client"],
   f'{c["persona"]}, {c["age"]}. ${c["portfolio"]:,}. Plan {pct(c["plan"])}. Band ±{round(c["band"]*100)} pts. Risk {c["risk"]}, capacity {c["capacity"]}. Goal: {c["goal"]["name"]} ({c["goal"]["years"]} yrs, target ${c["goal"]["target"]:,}, {round(c["goal"]["earmark"]*100)}% of portfolio). Restrictions: {c["restrictions"]}. Contact: {c["contact_text"]}. Consent: {"yes" if c["consent"] else "NO"}.',
   s["event_text"] + (f' Note text: "{s["event"]["note_text"]}"' if s["event"].get("note_text") else ""), s["source"],
   e["tier"], ", ".join(sorted(e["flags"])) or "-", e["message"], e["goal_after"] or "-",
   " • ".join(s["must_include"]), " • ".join(s["must_not"]), s["must_not_trade"],
   json.dumps(sig) if sig else "-", s["tests"],
   r["tier"], r["message"], r.get("max_drift_pts", "-"), r.get("band_pts", "-"),
   f'{r.get("goal_before") or "n/a"} > {r.get("goal_after") or "n/a"}', t, "PASS" if ok else "FAIL"]
  for s, c, r, e, ok, t, sig in rows],
 [7, 22, 34, 8, 60, 50, 16, 9, 30, 12, 12, 48, 38, 12, 34, 30, 8, 10, 9, 8, 18, 36, 7])

sheet("Clients",
 ["Client", "Persona", "Age", "Portfolio ($)", "Equity %", "Fixed income %", "Cash %", "Alternatives %", "Single stock %", "Band (pts)", "Risk tolerance", "Risk capacity",
  "Goal", "Years", "Share of portfolio funding goal", "Target ($)", "Restrictions", "Contact", "Monitoring consent"],
 [[k, c["persona"], c["age"], c["portfolio"]] + [round(c["plan"][x] * 100) for x in CLASSES] + [round(c["band"] * 100), c["risk"], c["capacity"],
   c["goal"]["name"], c["goal"]["years"], f'{round(c["goal"]["earmark"]*100)}%', c["goal"]["target"], c["restrictions"], c["contact_text"], "Yes" if c["consent"] else "No"]
  for k, c in CLIENTS.items()], [8, 30, 8, 13, 8, 9, 8, 10, 9, 8, 14, 10, 30, 7, 12, 12, 50, 32, 10])

sheet("Tier framework", ["Factor", "Question", "Level 1", "Level 2", "Level 3"], [
 ["Size", "How much of the portfolio does the proposed adjustment or confirmed cash flow move? (design decision, see About)", "Under 3%", "3% to 10%", "Over 10%"],
 ["Risk shift", "How far does the mix move from plan?", "Inside the client's band", "Outside the band, up to twice it", "More than twice the band, or a mandate cap is breached"],
 ["Goal impact", "Is a goal touched?", "Every goal on track (25th percentile meets target)", "A goal at risk (median meets target, 25th percentile does not)", "A goal off track, or at risk and due within 2 years"],
 ["Client sensitivity", "Is something delicate involved?", "Nothing sensitive", "Confirmed life event; concentration above limit; locked holding limits action; exclusion or restricted-list product blocked; profile conflict", "Sensitive personal situation; mandate limit breached"],
 ["Reversibility", "Can it be undone?", "Reversible", "Hard to reverse (locked or illiquid)", "Cannot be undone"],
 ["", "", "", "", ""],
 ["Tier", "Meaning", "Approval", "Client message", ""],
 ["Tier 1: Inform", "Within range, nothing to decide", "None; daily digest", "None unless RM chooses", ""],
 ["Tier 2: Review", "A proposed adjustment or a conversation", "RM approves, edits or rejects", "Drafted; released only by the RM (held if client prefers calls or no automated messages)", ""],
 ["Tier 3: Significant", "Large, goal-critical, sensitive or hard to reverse", "RM plus a second approver, after a client conversation. Never automated.", "Held until the RM has spoken to the client", ""],
 ["Paused", "A control stopped processing: no consent, stale/implausible/missing/conflicting data, currency mismatch, hidden instruction", "Nothing released; RM notified with the reason", "None", ""],
], [20, 60, 40, 55, 45])

sheet("Checks", ["Check", "What it tests", "On failure", "Pass bar", "Covered by"], [
 ["Monitoring consent", "Client has consented before any analysis", "Paused; consent task raised", "100%", "S-29"],
 ["Data freshness", "Feed age within limit", "Paused with data-age notice", "100%", "S-30"],
 ["Data plausibility", "No single-update move beyond limit", "Held for verification", "100%", "S-33"],
 ["Data completeness", "Every held instrument has a price", "Paused: 'Insufficient evidence to assess'", "100%", "S-34"],
 ["Source agreement", "Sources do not contradict each other", "Paused: 'Sources disagree'", "100%", "S-35"],
 ["Currency / units", "Prices normalised to USD", "Paused: data unavailable in USD", "100%", "S-36"],
 ["Untrusted-text scan", "Notes and news are data, never instructions", "Quarantined; compliance alerted; Paused", "100%", "S-32"],
 ["Life-event certainty and attribution", "'Might' is not 'did'; relative's event is not the client's; RM confirms before profile changes", "Recorded as possible / not attributed", "100%", "S-10, S-13, S-18, S-19"],
 ["Restrictions and preferences", "No trade touches an untradable holding, no-buy class or exclusion", "Candidate blocked; alternative proposed", "0 violations", "S-05, S-08, S-20, S-21, S-23"],
 ["Mandate limits", "Post-event weights within mandate", "Escalated to compliance", "0 unreported", "S-28"],
 ["Restricted list", "No restricted product proposed", "Blocked and logged", "0 violations", "S-27"],
 ["Concentration limit", "Single holding vs policy limit, with documented exception", "Flagged; trim unless excepted", "Flag every time", "S-05, S-06, S-21"],
 ["Profile consistency", "CRM risk tolerance agrees with plan", "No proposal until RM reconciles", "100%", "S-31"],
 ["Prohibited language", "No guarantee/promise/forecast wording in any draft", "Draft blocked", "0 findings", "Draft-text tests (separate)"],
 ["Numbers traceable", "Every figure comes from an engine, the client record or policy", "Draft blocked", "0 findings", "Draft-text tests (separate)"],
 ["Human approval gate", "Tier sets who must approve before release", "Nothing released", "100%", "Workflow tests (separate)"],
], [32, 62, 42, 14, 28])

sheet("Assumptions", ["Parameter", "Value", "Unit", "Note"],
 [[f"Expected annual return: {CLASS_NAMES[c]}", ASSUMPTIONS["return"][c], "per year", "Used only to compare goal targets with a projected range. Not a forecast."] for c in CLASSES]
 + [[f"Annual volatility: {CLASS_NAMES[c]}", ASSUMPTIONS["vol"][c], "per year", "Illustrative."] for c in CLASSES]
 + [["Correlation between classes", ASSUMPTIONS["correlation"], "ratio", "One flat value for every pair. A simplification."],
    ["25th-percentile z-score", ASSUMPTIONS["z25"], "", "Lognormal approximation."],
    ["Concentration limit (single holding)", ASSUMPTIONS["concentration_limit"], "share of portfolio", "Placeholder policy limit."],
    ["Maximum market-data age", ASSUMPTIONS["max_feed_age_min"], "minutes", "Placeholder."],
    ["Bad tick: diversified class", ASSUMPTIONS["bad_tick"]["diversified"], "move in one update", "Move beyond this is treated as implausible."],
    ["Bad tick: single stock", ASSUMPTIONS["bad_tick"]["single_stock"], "move in one update", ""],
    ["Size level 2 starts at", ASSUMPTIONS["size_l2"], "share of portfolio", "From the concept document."],
    ["Size level 3 starts above", ASSUMPTIONS["size_l3"], "share of portfolio", "From the concept document."],
    ["Goal level 3 if at risk and due within", ASSUMPTIONS["goal_l3_years"], "years", "From the concept document."],
    ["Goal status rule", "On track: 25th percentile meets target. At risk: median does, 25th does not. Off track: median misses.", "", "From the concept document."]],
 [44, 40, 20, 70])

sheet("Review (to complete)", ["ID", "Advisor reviewer", "Compliance reviewer", "Agree with expected tier? (Y/N)", "Comments / corrected expectation", "Date"],
      [[s["id"], "", "", "", "", ""] for s in S], [8, 24, 24, 20, 60, 12])

out = os.path.join(HERE, "DWN_Golden_Data_Set_v2.xlsx")
wb.save(out)
print("wrote", out, "and golden_set.json;", len(S), "scenarios;", dict(tiers))
